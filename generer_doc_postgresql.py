"""
Documentation Word - Guide PostgreSQL pour le projet sentiment_etudiants
"""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import datetime

doc = Document()

# ── Marges ─────────────────────────────────────────────────────────────────
for section in doc.sections:
    section.top_margin    = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin   = Cm(3)
    section.right_margin  = Cm(2.5)

# ── Couleurs ────────────────────────────────────────────────────────────────
BLEU      = RGBColor(0x1F, 0x49, 0x7D)
BLEU2     = RGBColor(0x2E, 0x86, 0xC1)
VERT      = RGBColor(0x1F, 0x8A, 0x70)
ROUGE     = RGBColor(0xC0, 0x39, 0x2B)
ORANGE    = RGBColor(0xD3, 0x54, 0x00)
GRIS      = RGBColor(0x44, 0x44, 0x44)

def set_cell_bg(cell, hex_color):
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd  = OxmlElement('w:shd')
    shd.set(qn('w:val'),   'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'),  hex_color)
    tcPr.append(shd)

def titre(texte):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(texte)
    r.bold = True
    r.font.size = Pt(22)
    r.font.color.rgb = BLEU

def h1(texte):
    p = doc.add_heading(level=1)
    p.clear()
    r = p.add_run(texte)
    r.bold = True
    r.font.size = Pt(15)
    r.font.color.rgb = BLEU

def h2(texte):
    p = doc.add_heading(level=2)
    p.clear()
    r = p.add_run(texte)
    r.bold = True
    r.font.size = Pt(12)
    r.font.color.rgb = BLEU2

def para(texte, bold=False, couleur=None):
    p = doc.add_paragraph()
    r = p.add_run(texte)
    r.bold = bold
    r.font.size = Pt(11)
    if couleur:
        r.font.color.rgb = couleur
    return p

def etape(numero, titre_etape, commande, explication, resultat_attendu=None):
    """Bloc visuel pour une étape numérotée."""
    # Numéro + titre
    p = doc.add_paragraph()
    r1 = p.add_run(f"Étape {numero} — ")
    r1.bold = True
    r1.font.size = Pt(12)
    r1.font.color.rgb = BLEU
    r2 = p.add_run(titre_etape)
    r2.bold = True
    r2.font.size = Pt(12)
    r2.font.color.rgb = GRIS

    # Commande
    pc = doc.add_paragraph()
    pc.paragraph_format.left_indent = Cm(1)
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), 'EBF5FB')
    pc._p.get_or_add_pPr().append(shd)
    rc = pc.add_run(f"  {commande}  ")
    rc.font.name  = 'Courier New'
    rc.font.size  = Pt(11)
    rc.font.color.rgb = RGBColor(0x17, 0x40, 0x6B)
    rc.bold = True

    # Explication
    pe = doc.add_paragraph()
    pe.paragraph_format.left_indent = Cm(1)
    re = pe.add_run(f"→  {explication}")
    re.font.size = Pt(10.5)
    re.font.color.rgb = GRIS

    # Résultat attendu
    if resultat_attendu:
        pr = doc.add_paragraph()
        pr.paragraph_format.left_indent = Cm(1)
        shd2 = OxmlElement('w:shd')
        shd2.set(qn('w:val'), 'clear')
        shd2.set(qn('w:color'), 'auto')
        shd2.set(qn('w:fill'), 'EAFAF1')
        pr._p.get_or_add_pPr().append(shd2)
        rr = pr.add_run(f"  ✅  Résultat attendu : {resultat_attendu}")
        rr.font.size = Pt(10)
        rr.font.color.rgb = VERT

    doc.add_paragraph()

def note(texte, emoji="💡"):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent  = Cm(0.8)
    p.paragraph_format.right_indent = Cm(0.8)
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'),   'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'),  'FEF9E7')
    p._p.get_or_add_pPr().append(shd)
    r = p.add_run(f"  {emoji}  {texte}")
    r.font.size = Pt(10.5)
    r.font.color.rgb = RGBColor(0x7D, 0x60, 0x08)
    doc.add_paragraph()

def code(texte):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(1)
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), 'F2F3F4')
    p._p.get_or_add_pPr().append(shd)
    r = p.add_run(f"  {texte}")
    r.font.name  = 'Courier New'
    r.font.size  = Pt(10)
    r.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)
    return p

def separateur():
    p = doc.add_paragraph()
    p.add_run('─' * 72).font.color.rgb = RGBColor(0xCC, 0xCC, 0xCC)


# ══════════════════════════════════════════════════════════════════════════════
# PAGE DE TITRE
# ══════════════════════════════════════════════════════════════════════════════
doc.add_paragraph()
doc.add_paragraph()
titre("🐘  GUIDE PostgreSQL")
doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Connexion · Déconnexion · Statut · Commandes essentielles")
r.font.size = Pt(13)
r.font.color.rgb = GRIS
separateur()
doc.add_paragraph()
p2 = doc.add_paragraph()
p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
p2.add_run(f"Projet : Analyse de Sentiments des Commentaires Étudiants\n").font.size = Pt(11)
p2.add_run(f"Date : {datetime.date.today().strftime('%d/%m/%Y')}").font.size = Pt(11)
doc.add_page_break()

# ══════════════════════════════════════════════════════════════════════════════
# SECTION 1 — INFORMATIONS DE CONNEXION
# ══════════════════════════════════════════════════════════════════════════════
h1("1.  Informations de connexion à retenir")
para("Voici les identifiants de ta base de données PostgreSQL pour ce projet :", bold=True)
doc.add_paragraph()

lignes  = [
    ('Nom de la base de données', 'sentiment_etudiants'),
    ('Utilisateur (username)',    'postgres'),
    ('Mot de passe (password)',   'admin123'),
    ('Numéro de port',            '5432'),
    ('Hôte (host)',               'localhost'),
]
table = doc.add_table(rows=1 + len(lignes), cols=2)
table.style = 'Table Grid'
table.alignment = WD_TABLE_ALIGNMENT.CENTER

# En-tête
set_cell_bg(table.rows[0].cells[0], '1F497D')
set_cell_bg(table.rows[0].cells[1], '1F497D')
for j, txt in enumerate(['Paramètre', 'Valeur']):
    r = table.rows[0].cells[j].paragraphs[0].add_run(txt)
    r.bold = True
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    r.font.size = Pt(11)

for i, (param, val) in enumerate(lignes):
    bg = 'EAF2FF' if i % 2 == 0 else 'FFFFFF'
    set_cell_bg(table.rows[i+1].cells[0], bg)
    set_cell_bg(table.rows[i+1].cells[1], bg)
    table.rows[i+1].cells[0].paragraphs[0].add_run(param).font.size = Pt(11)
    rv = table.rows[i+1].cells[1].paragraphs[0].add_run(val)
    rv.bold = True
    rv.font.size = Pt(11)
    rv.font.color.rgb = BLEU

doc.add_paragraph()
note(
    "Garde ces informations en sécurité. Ne les partage pas publiquement "
    "(surtout pas sur GitHub).",
    "🔒"
)

doc.add_page_break()

# ══════════════════════════════════════════════════════════════════════════════
# SECTION 2 — VÉRIFIER LE STATUT
# ══════════════════════════════════════════════════════════════════════════════
h1("2.  Vérifier l'état de PostgreSQL")
para("Avant toute chose, vérifie si PostgreSQL est démarré ou non.")
doc.add_paragraph()

etape(
    1,
    "Ouvrir un terminal",
    "Ctrl + Alt + T",
    "Raccourci clavier pour ouvrir un nouveau terminal sous Ubuntu."
)

etape(
    2,
    "Vérifier le statut",
    "sudo service postgresql status",
    "Cette commande affiche si PostgreSQL est actif ou arrêté. "
    "Ubuntu te demandera ton mot de passe système (pas le mot de passe postgres).",
    "online  ou  active (running)"
)

para("Tu verras l'une de ces deux réponses :", bold=True)
doc.add_paragraph()

table2 = doc.add_table(rows=3, cols=2)
table2.style = 'Table Grid'
table2.alignment = WD_TABLE_ALIGNMENT.CENTER
set_cell_bg(table2.rows[0].cells[0], '1F497D')
set_cell_bg(table2.rows[0].cells[1], '1F497D')
for j, txt in enumerate(['Message affiché', 'Signification']):
    r = table2.rows[0].cells[j].paragraphs[0].add_run(txt)
    r.bold = True
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

set_cell_bg(table2.rows[1].cells[0], 'EAFAF1')
set_cell_bg(table2.rows[1].cells[1], 'EAFAF1')
table2.rows[1].cells[0].paragraphs[0].add_run('active (running)  ✅').font.size = Pt(10)
table2.rows[1].cells[1].paragraphs[0].add_run('PostgreSQL est démarré, prêt à recevoir des connexions.').font.size = Pt(10)

set_cell_bg(table2.rows[2].cells[0], 'FDEDEC')
set_cell_bg(table2.rows[2].cells[1], 'FDEDEC')
table2.rows[2].cells[0].paragraphs[0].add_run('inactive (dead)  ❌').font.size = Pt(10)
table2.rows[2].cells[1].paragraphs[0].add_run('PostgreSQL est arrêté. Il faut le démarrer (voir section 3).').font.size = Pt(10)

doc.add_paragraph()
doc.add_page_break()

# ══════════════════════════════════════════════════════════════════════════════
# SECTION 3 — DÉMARRER / ARRÊTER / REDÉMARRER
# ══════════════════════════════════════════════════════════════════════════════
h1("3.  Démarrer, Arrêter et Redémarrer PostgreSQL")
doc.add_paragraph()

etape(
    1,
    "Démarrer PostgreSQL",
    "sudo service postgresql start",
    "Lance le serveur PostgreSQL. À faire au début de chaque session de travail "
    "si PostgreSQL n'est pas déjà actif.",
    "* Starting PostgreSQL 16 database server"
)

etape(
    2,
    "Arrêter PostgreSQL",
    "sudo service postgresql stop",
    "Arrête proprement le serveur. À faire quand tu as fini de travailler "
    "et que tu n'as plus besoin de la base de données.",
    "* Stopping PostgreSQL 16 database server"
)

etape(
    3,
    "Redémarrer PostgreSQL",
    "sudo service postgresql restart",
    "Arrête puis redémarre le serveur. Utile après une modification de configuration "
    "ou si le serveur ne répond plus correctement.",
    "* Restarting PostgreSQL 16 database server"
)

note(
    "Pour ce projet, tu dois démarrer PostgreSQL AVANT de lancer l'API FastAPI. "
    "Sans PostgreSQL actif, FastAPI ne peut pas sauvegarder les commentaires.",
    "⚠️"
)

doc.add_page_break()

# ══════════════════════════════════════════════════════════════════════════════
# SECTION 4 — SE CONNECTER À LA BASE DE DONNÉES
# ══════════════════════════════════════════════════════════════════════════════
h1("4.  Se connecter à la base de données")
para("Pour entrer dans PostgreSQL et exécuter des commandes SQL directement.")
doc.add_paragraph()

etape(
    1,
    "Ouvrir un terminal",
    "Ctrl + Alt + T",
    "Ouvre un nouveau terminal."
)

etape(
    2,
    "Se connecter",
    "psql -U postgres -d sentiment_etudiants",
    "Cette commande se connecte à la base 'sentiment_etudiants' avec l'utilisateur 'postgres'.\n"
    "   -U postgres         → nom d'utilisateur\n"
    "   -d sentiment_etudiants  → nom de la base de données",
    "sentiment_etudiants=#   (le curseur change, tu es connecté)"
)

note(
    "Si PostgreSQL te demande un mot de passe, tape :  admin123  puis Entrée. "
    "Le mot de passe n'apparaît pas à l'écran pendant la saisie, c'est normal.",
    "🔑"
)

doc.add_page_break()

# ══════════════════════════════════════════════════════════════════════════════
# SECTION 5 — COMMANDES SQL UTILES
# ══════════════════════════════════════════════════════════════════════════════
h1("5.  Commandes utiles une fois connecté")
para("Une fois que tu vois le curseur  sentiment_etudiants=#  tu peux taper :")
doc.add_paragraph()

table3 = doc.add_table(rows=8, cols=2)
table3.style = 'Table Grid'
table3.alignment = WD_TABLE_ALIGNMENT.CENTER
set_cell_bg(table3.rows[0].cells[0], '1F497D')
set_cell_bg(table3.rows[0].cells[1], '1F497D')
for j, txt in enumerate(['Commande', 'Ce que ça fait']):
    r = table3.rows[0].cells[j].paragraphs[0].add_run(txt)
    r.bold = True
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

cmds = [
    (r'\dt',                         'Afficher toutes les tables de la base'),
    (r'\l',                          'Lister toutes les bases de données'),
    (r'\du',                         'Lister tous les utilisateurs'),
    ('SELECT * FROM commentaires;',  'Afficher tous les commentaires sauvegardés'),
    ('SELECT COUNT(*) FROM commentaires;', 'Compter le nombre de commentaires'),
    ('SELECT * FROM commentaires ORDER BY date_creation DESC LIMIT 5;',
     'Voir les 5 derniers commentaires'),
    (r'\q',                          'Quitter et se déconnecter de PostgreSQL'),
]
for i, (cmd, desc) in enumerate(cmds):
    bg = 'EAF2FF' if i % 2 == 0 else 'FFFFFF'
    set_cell_bg(table3.rows[i+1].cells[0], bg)
    set_cell_bg(table3.rows[i+1].cells[1], bg)
    r1 = table3.rows[i+1].cells[0].paragraphs[0].add_run(cmd)
    r1.font.name  = 'Courier New'
    r1.font.size  = Pt(10)
    r1.bold       = True
    r1.font.color.rgb = RGBColor(0x17, 0x40, 0x6B)
    table3.rows[i+1].cells[1].paragraphs[0].add_run(desc).font.size = Pt(10)

doc.add_paragraph()
note(
    "Toutes les commandes SQL (SELECT, INSERT, etc.) doivent se terminer par un "
    "point-virgule  ;  sinon PostgreSQL attend la suite.",
    "💡"
)

doc.add_page_break()

# ══════════════════════════════════════════════════════════════════════════════
# SECTION 6 — SE DÉCONNECTER
# ══════════════════════════════════════════════════════════════════════════════
h1("6.  Se déconnecter de PostgreSQL")
doc.add_paragraph()

etape(
    1,
    "Quitter la console PostgreSQL",
    r"\q",
    "Tape \\q puis Entrée. Cela ferme la session psql et te ramène "
    "au terminal normal.",
    "Tu retrouves le prompt normal :  nasser@pc:~$"
)

note(
    "Se déconnecter de la console psql (\\q) ne signifie pas qu'on arrête PostgreSQL. "
    "Le serveur continue à tourner en arrière-plan. "
    "Pour arrêter le serveur, utilise :  sudo service postgresql stop",
    "⚠️"
)

doc.add_page_break()

# ══════════════════════════════════════════════════════════════════════════════
# SECTION 7 — ORDRE DE LANCEMENT DU PROJET
# ══════════════════════════════════════════════════════════════════════════════
h1("7.  Ordre de lancement du projet complet")
para("À chaque fois que tu veux travailler sur le projet, suis cet ordre :")
doc.add_paragraph()

etapes_lancement = [
    ("Terminal 1",
     "sudo service postgresql start",
     "Démarrer la base de données",
     "active (running)"),
    ("Terminal 1",
     "cd ~/Documents/sentiment_etudiants && source venv/bin/activate && uvicorn backend.main:app --reload --host 0.0.0.0 --port 8000",
     "Lancer le serveur FastAPI (API)",
     "Application startup complete."),
    ("Terminal 2",
     "cd ~/Documents/sentiment_etudiants && source venv/bin/activate && streamlit run dashboard/app.py",
     "Lancer le Dashboard Streamlit",
     "Onglet navigateur ouvert sur http://localhost:8501"),
    ("Android Studio",
     "Démarrer l'émulateur puis : flutter run",
     "Lancer l'application mobile Flutter",
     "Application affichée sur l'émulateur"),
]

for i, (terminal, cmd, desc, res) in enumerate(etapes_lancement):
    p = doc.add_paragraph()
    r1 = p.add_run(f"  {i+1}.  [{terminal}]  ")
    r1.bold = True
    r1.font.color.rgb = BLEU
    r1.font.size = Pt(11)

    pc = doc.add_paragraph()
    pc.paragraph_format.left_indent = Cm(1.5)
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), 'EBF5FB')
    pc._p.get_or_add_pPr().append(shd)
    rc = pc.add_run(f"  {cmd}  ")
    rc.font.name  = 'Courier New'
    rc.font.size  = Pt(9.5)
    rc.font.color.rgb = RGBColor(0x17, 0x40, 0x6B)

    pd = doc.add_paragraph()
    pd.paragraph_format.left_indent = Cm(1.5)
    pd.add_run(f"→  {desc}").font.size = Pt(10)

    pr = doc.add_paragraph()
    pr.paragraph_format.left_indent = Cm(1.5)
    rr = pr.add_run(f"✅  {res}")
    rr.font.size = Pt(10)
    rr.font.color.rgb = VERT
    doc.add_paragraph()

separateur()
doc.add_paragraph()
p_fin = doc.add_paragraph()
p_fin.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_fin = p_fin.add_run("🐘  PostgreSQL est le moteur silencieux de ton projet — il tourne en arrière-plan\net garde la mémoire de tous les commentaires analysés.")
r_fin.bold = True
r_fin.font.size = Pt(11)
r_fin.font.color.rgb = BLEU

# ── Sauvegarde ─────────────────────────────────────────────────────────────
chemin = "/home/nasser/Documents/sentiment_etudiants/Guide_PostgreSQL.docx"
doc.save(chemin)
print(f"✅  Document généré : {chemin}")
